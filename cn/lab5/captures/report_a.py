from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SH = "../shots/"
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.8)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

def para(text="", size=11, bold=False, center=False, after=6):
    p = doc.add_paragraph()
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    for i, t in enumerate(text.split("**")):
        if not t: continue
        r = p.add_run(t); r.font.size = Pt(size); r.bold = bold or (i % 2 == 1)
    return p

def head(t): para(t, 14, True, after=4).paragraph_format.keep_with_next = True
def sub(t): para(t, 11, True, after=2).paragraph_format.keep_with_next = True
def img(f, w=6.4, maxh=5.6):
    W, H = Image.open(SH + f).size; w = min(w, maxh * W / H)
    doc.add_picture(SH + f, width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
def cap(t): para(t, 9, center=True, after=10)
def bullet(t):
    p = para(t); p.style = doc.styles["List Bullet"]

def table(rows, widths=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = ""
            run = c.paragraphs[0].add_run(v); run.font.size = Pt(10); run.bold = header and i == 0
            if widths: c.width = Inches(widths[j])
            c.paragraphs[0].paragraph_format.keep_with_next = i < len(rows) - 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ------------------------------------------------------------------ header
para("Computer Networks (EE353)", 15, True, True, 2)
para("Lab 5: TCP Traffic Analysis using Wireshark", 12, True, True, 2)
para("Arqam Waheed 553514       BSDS 3A        8 October 2026", 11, center=True, after=12)

head("Setup")
para("Wireshark 3.6.2 on Ubuntu, Wi-Fi interface wlp0s20f3 (my IP 192.168.18.55). "
     "Capture file lab5_tcp.pcapng (1519 packets, 1511 TCP). During the capture I opened example.com over HTTP and HTTPS "
     "with curl, then tried port 81 on my router to get a reset.")
img("a0_terminal.png", 5.6)
cap("Commands used to generate TCP traffic")

# ------------------------------------------------------------------ task 1
head("Task 1) Identify TCP Traffic")
img("a1_tcp.png")
cap("Filter: tcp, frame 201")
table([["Field", "Observation"],
       ["Source IP", "192.168.18.55"],
       ["Destination IP", "104.20.23.154 (example.com)"],
       ["Source Port", "50696"],
       ["Destination Port", "80"],
       ["Sequence Number", "1 (raw 3940007916)"],
       ["Acknowledgment Number", "1 (raw 2858064501)"],
       ["Flags", "0x018 (PSH, ACK)"],
       ["Window Size", "502 × 128 = 64256 bytes"],
       ["TCP Header Length", "32 bytes (20 + 12 bytes of options)"],
       ["Checksum / Status", "0x52ff [unverified]"],
       ["Higher Layer Protocol", "HTTP (GET / HTTP/1.1)"]], [2.2, 4.4])

# ------------------------------------------------------------------ task 2
head("Task 2) Find the Three-Way Handshake")
img("a2_handshake.png")
cap("Filter: tcp.stream eq 25, frames 198, 199, 200")
table([["Packet", "Frame", "Ports", "Seq", "Ack"],
       ["SYN", "198", "50696 → 80", "0 (raw 3940007915)", "—"],
       ["SYN/ACK", "199", "80 → 50696", "0 (raw 2858064500)", "1 (raw 3940007916)"],
       ["ACK", "200", "50696 → 80", "1", "1 (raw 2858064501)"]], [1.0, 0.7, 1.3, 1.8, 1.8])
para("**Purpose of SYN:** It asks to open a connection and sends the client's starting sequence number (3940007915).")
para("**What SYN/ACK acknowledges:** The client's SYN. Its Ack is the client's ISN + 1 (3940007916). It also carries the server's own ISN.")
para("**Why the final ACK:** It confirms the server's SYN, so both sides know the other can send and receive. Then data can flow.")
para("**Who initiated:** My laptop (192.168.18.55), because it sent the SYN.")

# ------------------------------------------------------------------ task 3
head("Task 3) Sequence and Acknowledgment Analysis")
img("a3_seqack.png")
cap("Filter: tcp.stream eq 25, ACK in frame 202 for data in frame 201")
table([["Field", "Observation"],
       ["Data packet", "Frame 201 (HTTP GET), client → server"],
       ["Sequence number", "1"],
       ["TCP segment length", "75 bytes"],
       ["Next sequence number", "76 (1 + 75)"],
       ["ACK packet", "Frame 202, server → client"],
       ["Acknowledgment number", "76"],
       ["Time between data and ACK", "2.021087 − 2.016559 = 0.004528 s (4.5 ms)"]], [2.2, 4.4])
para("**How seq/ack give reliable, ordered delivery:** Each byte has a sequence number, so the receiver can put segments "
     "back in order and spot gaps. The ACK number tells the sender the next byte expected (76 here means all 75 bytes arrived). "
     "Anything not acknowledged in time is resent. The SYN used up sequence number 0, which is why data starts at 1.")

# ------------------------------------------------------------------ task 4
head("Task 4) Analyze TCP Flags")
table([["Filter", "Packets", "Meaning"],
       ["tcp.flags.syn == 1", "13", "Opening a connection (SYN and SYN/ACK)"],
       ["tcp.flags.ack == 1", "1501", "Ack number is valid; set on almost every packet after the SYN"],
       ["tcp.flags.fin == 1", "11", "Sender has no more data and wants to close"],
       ["tcp.flags.reset == 1", "5", "Connection aborted immediately"]], [1.8, 0.8, 4.0])
img("a4_rst.png")
cap("Filter: tcp.flags.reset == 1, frame 717")
para("Frame 717 is the router (192.168.18.1) answering my SYN to port 81 with RST, ACK because nothing listens on that port.")

# ------------------------------------------------------------------ task 5
head("Task 5) Follow TCP Stream")
img("a5_follow.png", 5.8)
cap("Analyze → Follow → TCP Stream on frame 201")
para("Stream 25. Client 192.168.18.55:50696, server 104.20.23.154:80. Wireshark applied tcp.stream eq 25, showing only "
     "the 12 packets of this connection. Red is what the client sent (GET request), blue is the server's reply (HTTP 200 OK).")
para("**What a TCP stream is:** One connection, identified by the two IPs and two ports.")
para("**Why it is useful:** It hides unrelated traffic so you can read one conversation from start to end.")
para("**Why HTTPS data is not readable:** TLS encrypts the payload. Following the HTTPS stream (34) only shows random bytes.")

# ------------------------------------------------------------------ task 6
head("Task 6) Observe Connection Termination")
img("a6_fin.png")
cap("Filter: tcp.flags.fin == 1, frame 207")
table([["Frame", "Direction", "Flags", "Seq / Ack"],
       ["207", "client → server", "FIN, ACK", "Seq 76, Ack 937"],
       ["225", "server → client", "FIN, ACK", "Seq 937, Ack 77"],
       ["226", "client → server", "ACK", "Seq 77, Ack 938"]], [0.8, 1.8, 1.4, 2.6])
para("My client closed first. The server acknowledged the FIN and sent its own FIN in the same packet, so only 3 packets were needed. "
     "The HTTPS connection (frames 368–371) closed with the full 4 packets.")
para("RST was also seen: frame 717 (closed port 81 on the router) and frames 887/905 where the Ubuntu connectivity-check server reset its connection.")
para("**Why multiple packets in both directions?** TCP is full duplex, so each side closes its own direction with a FIN and the other side has to ACK it.")

# ------------------------------------------------------------------ task 7
head("Task 7) Analyze TCP Port 443")
img("a7_443.png")
cap("Filter: tcp.port == 443, frame 342")
para("1482 packets matched. Stream 34 (192.168.18.55:43358 → 104.20.23.154:443) did a normal handshake: SYN (342), SYN/ACK (343), ACK (344), "
     "then a TLS Client Hello in frame 349. Port 443 is normally HTTPS, but the port alone does not prove it. Wireshark only knows it is "
     "TLS because it dissected the payload.")

# ------------------------------------------------------------------ task 8
head("Task 8) Compare TCP and UDP")
table([["Feature", "TCP", "UDP"],
       ["Connection setup", "Yes (3-way handshake)", "No"],
       ["Reliability", "Yes (ACKs, retransmission)", "No"],
       ["Ordering", "Yes (sequence numbers)", "No"],
       ["Acknowledgments", "Yes", "No"],
       ["Flow control", "Yes (receive window)", "No"],
       ["Minimum transport header", "20 bytes", "8 bytes"],
       ["Examples", "HTTP, HTTPS, SSH, email", "DNS, DHCP, QUIC, video"]], [2.2, 2.2, 2.2])
bullet("Here TCP needed 3 packets (198–200) before sending any data. In Lab 4 the DNS lookup was just a query and a reply.")
bullet("The TCP header in my packets was 32 bytes (20 + 12 bytes of timestamp options). The UDP header in Lab 4 was always 8 bytes.")
bullet("Frame 202 acknowledged the GET in frame 201. UDP had nothing like this.")

# ------------------------------------------------------------------ viva
head("Viva Questions")
qa = [("What is TCP and why is it connection-oriented?", "A reliable transport protocol. It sets up a connection with a handshake before sending data."),
      ("Three-way handshake?", "Client sends SYN, server replies SYN/ACK, client sends ACK."),
      ("Purpose of SYN?", "To start a connection and share the initial sequence number."),
      ("Sequence vs acknowledgment number?", "Seq is the number of the first byte in this segment. Ack is the next byte the sender expects to receive."),
      ("Why does TCP use ACKs?", "So the sender knows what arrived and can resend what didn't."),
      ("What is the receive window?", "How many bytes the receiver can still accept. The sender cannot have more than this unacknowledged."),
      ("Flow control vs congestion control?", "Flow control protects the receiver (window). Congestion control protects the network (slows down on loss)."),
      ("FIN vs RST?", "FIN closes gracefully after all data is sent. RST aborts the connection immediately."),
      ("What does Follow TCP Stream do?", "Filters one connection (tcp.stream eq N) and shows its payload as a conversation."),
      ("Why might a website not produce TCP/443?", "The browser may use HTTP/3 (QUIC over UDP/443), or the page was cached."),
      ("What is a TCP segment?", "The TCP header plus the data it carries."),
      ("How does TCP differ from UDP?", "TCP is connection-oriented, reliable and ordered. UDP is connectionless with no guarantees."),
      ("Retransmission / duplicate ACK?", "Packet loss or reordering on the path. Several duplicate ACKs make the sender resend early.")]
for i, (q, a) in enumerate(qa, 1):
    para(f"**{i}. {q}** {a}", after=3)

doc.save("../ArqamWaheed553514_Lab5_A.docx")
print("saved")
