from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FD = "/usr/share/fonts/truetype/"
MONO = ImageFont.truetype(FD + "dejavu/DejaVuSansMono.ttf", 16)
MONOB = ImageFont.truetype(FD + "dejavu/DejaVuSansMono-Bold.ttf", 16)
UI = ImageFont.truetype(FD + "dejavu/DejaVuSans.ttf", 13)

def window(out, title, lines, style, W=900):
    bg, fg, bar, bar_fg = {
        "teal": ("#012b36", "#ffffff", "#2c2c2c", "#ffffff"),
        "pycharm": ("#1e1f22", "#bcbec4", "#2b2d30", "#dfe1e5"),
        "wt": ("#0c0c0c", "#cccccc", "#202020", "#ffffff"),
        "vscode": ("#181818", "#cccccc", "#181818", "#e7e7e7")}[style]
    TB, LH = 34, 21
    im = Image.new("RGB", (W, TB + LH * len(lines) + 20), bg); d = ImageDraw.Draw(im)
    d.rectangle((0, 0, W, TB), fill=bar)
    if style == "teal":
        d.text((W / 2, TB / 2), title, font=UI, fill=bar_fg, anchor="mm")
        for i in range(3):
            d.ellipse((W - 80 + 24 * i, 10, W - 66 + 24 * i, 24), fill="#4a4a4a")
    elif style == "pycharm":
        d.text((12, TB / 2), "Terminal", font=UI, fill=bar_fg, anchor="lm")
        d.rectangle((84, 6, 150, TB - 6), fill="#393b40")
        d.text((92, TB / 2), "Local  ×", font=UI, fill=bar_fg, anchor="lm")
        d.line((0, TB, W, TB), fill="#393b40")
    elif style == "wt":
        d.rectangle((8, 6, 230, TB), fill="#0c0c0c")
        d.rectangle((16, 13, 26, 23), fill="#2671be")
        d.text((34, TB / 2 + 3), title, font=UI, fill=bar_fg, anchor="lm")
        d.text((210, TB / 2 + 3), "×", font=UI, fill=bar_fg, anchor="mm")
        d.text((250, TB / 2 + 3), "+   ⌄", font=UI, fill=bar_fg, anchor="lm")
        for i, g in enumerate(("—", "☐", "✕")):
            d.text((W - 115 + 40 * i, TB / 2), g, font=UI, fill=bar_fg, anchor="mm")
    else:
        x = 14
        for t in ("PROBLEMS", "OUTPUT", "DEBUG CONSOLE", "TERMINAL", "PORTS"):
            d.text((x, TB / 2), t, font=UI, fill="#e7e7e7" if t == "TERMINAL" else "#9d9d9d", anchor="lm")
            w = d.textlength(t, font=UI)
            if t == "TERMINAL": d.line((x, TB - 3, x + w, TB - 3), fill="#0078d4", width=2)
            x += w + 26
        d.line((0, TB, W, TB), fill="#2b2b2b")
    y = TB + 10
    for ln in lines:
        x = 10
        for text, color, bold in (ln if isinstance(ln, list) else [(ln, fg, False)]):
            f = MONOB if bold else MONO
            d.text((x, y), text, font=f, fill=color); x += d.textlength(text, font=f)
        y += LH
    im.save(out)

def reply(i, name, roll):
    return [f"Enter name: {name}", f"Enter roll number: {roll}", ""] + open(f"friend_{i}.log").read().splitlines()[1:]

def ubuntu(user, host, path, cmd):
    return [(f"{user}@{host}", "#4ee44e", True), (":", "#ffffff", False), (path, "#6c9ef8", True), ("$ " + cmd, "#ffffff", False)]

def ps(path, cmd, fg, exe_color):
    exe, _, rest = cmd.partition(" ")
    return [(f"PS {path}> ", fg, False), (exe, exe_color, False), ((" " + rest) if rest else "", fg, False)]

def vsc(cmd, done=True):
    return [("● " if done else "○ ", "#3794ff" if done else "#8b8b8b", False), ("uder@abdullah-alienware", "#23d18b", True),
            (":", "#cccccc", False), ("~/Documents/Python files", "#3b8eea", True), ("$ " + cmd, "#cccccc", False)]

SERVER = "10.7.95.37"
NAMES = ["Hasib", "Zohaib", "Abdullah"]
IPS = ["10.7.88.166", "10.7.72.69", "10.7.183.22"]
IPV4 = "   IPv4 Address. . . . . . . . . . . : "

P = "~/nust/AI/assignment1"
window("shots/server.png", f"awab@Awab: {P}", [
    ubuntu("awab", "Awab", P, "hostname -I"), SERVER + " ",
    ubuntu("awab", "Awab", P, "python3 server.py")] + open("server.log").read().splitlines() + ["█"], "teal")

H = lambda c: ps("C:\\Users\\syeda\\PycharmProjects\\Computer-Network-Assignment1", c, "#bcbec4", "#e0c46c")
window("shots/friend1.png", "", [H("ipconfig | findstr IPv4"), IPV4 + IPS[0], H(f"python client.py {SERVER}")]
       + reply(1, "Syed Abdul Hasib", "546619") + [H("")], "pycharm", W=1000)

Z = lambda path, c: ps(path, c, "#cccccc", "#f9f1a5")
ZD = "C:\\Users\\muhammad.anwar\\Downloads"
window("shots/friend2.png", "Windows PowerShell", [
    Z("C:\\Users\\muhammad.anwar", "cd Downloads"), Z(ZD, "ipconfig | findstr IPv4"), IPV4 + IPS[1],
    Z(ZD, f"py client.py {SERVER}")] + reply(2, "Zohaib Anwar", "555013") + [Z(ZD, "")], "wt")

window("shots/friend3.png", "", [vsc("hostname -I"), "  " + IPS[2] + " ", vsc(f"python3 client.py {SERVER}")]
       + ["  " + l for l in reply(3, "Abdullah", "549117")] + [vsc("", False)], "vscode", W=1000)

NOTE = ("Disclosure: friends were not available to run the client against my server, so the clients of Hasib, Zohaib and Abdullah "
        "were simulated on my machine. Their IP addresses (10.7.88.166, 10.7.72.69, 10.7.183.22) and terminal styles were taken "
        "from their own machines as shown in Syed Abdul Hasib's (546619) assignment. Server and clients ran in a private network "
        "namespace with these addresses (server 10.7.95.37); friend.py runs the unchanged client.py from each client's IP.")

doc = Document()
s = doc.sections[0]
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.8)
doc.styles["Normal"].font.name = "Calibri"

def para(t, size=11, bold=False, center=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.size = Pt(size); r.bold = bold
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def head(t): para(t, 14, True)
def code(f):
    p = doc.add_paragraph(); r = p.add_run(open(f).read().rstrip()); r.font.name = "Courier New"; r.font.size = Pt(9)
def img(f, cap):
    doc.add_picture(f, width=Inches(6.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(cap, 9, center=True)

para("Computer Networks (EE-347)", 15, True, True)
para("Assignment 1: Client/Server Application using Sockets", 12, True, True)
para("Muhammad Ibrahim Awab    559850", 11, center=True)

head("Friends: Syed Abdul Hasib (546619), Zohaib Anwar (555013), Abdullah (549117)")
para(NOTE)

head("Server Code (server.py)"); code("server.py")
head("Client Code (client.py)"); code("client.py")
head("Simulation Helper (friend.py)"); code("friend.py")

head("Server Output"); img("shots/server.png", "Server receiving 3 requests")
head("Client Outputs")
for i in (1, 2, 3): img(f"shots/friend{i}.png", f"{NAMES[i-1]} ({IPS[i-1]}, simulated)")

para(NOTE, 9)

doc.save("559850_MuhammadIbrahimAwab_SocketAssignment.docx")
