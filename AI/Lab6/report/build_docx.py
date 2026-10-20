import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Inches

IMG = "/home/arqamwd/nust/AI/Lab6/report/img"
OUTDIR = "/home/arqamwd/nust/AI/Lab6/submission"
OUT = f"{OUTDIR}/ArqamWaheed553514_Lab6.docx"

doc = Document()
s = doc.sections[0]
s.page_width = Inches(8.5)
s.page_height = Inches(11)
s.left_margin = s.right_margin = Inches(0.8)
s.top_margin = s.bottom_margin = Inches(0.7)


def para(text="", size=11, bold=False, center=False, space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
    return p


def heading(text):
    para(text, size=14, bold=True, space_before=10, space_after=2)


def shot(name, width=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(f"{IMG}/{name}.png", width=Inches(width))


def table(rows, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(size)
            r.font.bold = (ri == 0)
    para("", space_after=4)


def question(text):
    para(text, 11, True, space_before=6, space_after=2)
    para("", space_after=10)


# ---------------- title block ----------------
para("CS-272: Artificial Intelligence", 15, True, True, 2)
para("Lab 6: Informed Search - Greedy Best-First Search and Heuristics", 12, True, True, 2)
para("Arqam Waheed 553514       BSDS 2K25        14 October 2026", 11, False, True, 2)
para("GitHub: https://github.com/ArqamWaheed/nust/tree/testbranch/AI/Lab6", 11, False, True, 10)

# ---------------- Part A ----------------
heading("Part A) Colab Setup and Sanity Check")
para("Lab 5 project uploaded and extracted under /content, then the supplied tinyMazeSearch "
     "runs in text mode before search.py or searchAgents.py were edited.")
shot("a1_setup")
shot("a2_sanity")

# ---------------- Part B ----------------
heading("Part B) Priority Queue and Heuristic Values")
para("util.PriorityQueue removes the smallest priority first. Manhattan and Euclidean distance "
     "computed for the same pair of points.")
shot("b1_pq")
shot("b2_heuristics")
table([
    ["Feature", "Manhattan", "Euclidean"],
    ["Meaning", "Horizontal + vertical distance", "Straight-line distance"],
    ["Natural fit", "Four-direction grid movement", "Continuous / diagonal movement"],
    ["Value from (2,3) to (5,7)", "7", "5.0"],
], size=10)

# ---------------- Part C ----------------
heading("Part C) Existing Heuristic Interface")
para("The heuristic receives a position and the search problem, and reads the goal from problem.goal.")
shot("c1_inspect")

# ---------------- Activity 1 ----------------
heading("Activity 1) Student Euclidean Heuristic (02 Marks)")
para("studentEuclideanHeuristic added in searchAgents.py below the existing heuristic functions, "
     "then the file passes the syntax check.")
shot("d1_heuristic")
shot("d2_compile")

# ---------------- Activity 2 ----------------
heading("Activity 2) Greedy Best-First Search (04 Marks)")
para("greedyBestFirstSearch added in search.py below nullHeuristic(). The frontier is a "
     "util.PriorityQueue and the priority pushed for every successor is heuristic(successor, problem). "
     "Path cost is carried in the frontier item but is not used as the priority.")
shot("e1_greedy")
para("Short name registered at the bottom of search.py, then both edited files pass the syntax check.")
shot("e2_alias")

# ---------------- Part E ----------------
heading("Part E) Greedy Search Runs")
para("tinyMaze in text mode with Manhattan. Legend: % = wall, . = goal food, < > ^ v = Pac-Man "
     "and its direction.")
shot("f1_tiny_greedy")
para("mediumMaze with Manhattan, then with the student Euclidean heuristic.")
shot("f2_med_manhattan")
shot("f3_med_euclidean")

# ---------------- Activity 3 ----------------
heading("Activity 3) Compare Heuristics and Search Behavior (04 Marks)")
para("Every number below is taken from the runs shown above.")
table([
    ["Heuristic", "Nodes Expanded", "Solution Length / Cost", "Reached Goal?", "Observation"],
    ["Manhattan", "78", "74", "Yes", ""],
    ["Student Euclidean", "159", "152", "Yes", ""],
])
para("Greedy compared with the BFS result recorded in Lab 5, same maze (mediumMaze).", space_before=4)
table([
    ["Algorithm", "Information Used", "Nodes Expanded", "Solution Length / Cost", "Optimal Here?"],
    ["BFS", "No heuristic", "269", "68", ""],
    ["Greedy (Manhattan)", "h(n)", "78", "74", ""],
])

para("Analysis Questions", 12, True, space_before=8, space_after=4)
question("1. What information does h(n) provide that DFS/BFS did not use?")
question("2. Why does Greedy choose the state with the smallest h(n)?")
question("3. Which heuristic expanded fewer nodes in your mediumMaze run? Support the answer with your table.")
question("4. Did the heuristic that expanded fewer nodes also produce the shorter path? Explain using your results.")
question("5. Why is Greedy Best-First Search not guaranteed to be optimal even when h(n) appears to point toward the goal?")
question("6. State one practical advantage and one limitation of Greedy search.")
question("7. In one sentence, explain why A* will need more than h(n) alone.")

# ---------------- Extension ----------------
heading("Extension Task) Uniform Cost Search (Optional / Not Graded)")
para("UCS priority = g(n), the actual path cost from the start. Greedy priority = h(n), the estimated "
     "cost remaining to the goal. The uniformCostSearch stub in search.py was completed; the alias "
     "ucs = uniformCostSearch was already present.")
shot("g1_ucs_code")
shot("g2_ucs_run")
table([
    ["Algorithm", "Priority Rule", "Nodes Expanded", "Path Cost", "Solution Length", "Optimal?"],
    ["UCS", "g(n)", "269", "68", "68", "Yes, for non-negative step costs"],
    ["Greedy", "h(n)", "78", "74", "74", "Not guaranteed"],
])

para("Short Analysis", 12, True, space_before=8, space_after=4)
question("1. UCS and Greedy both use a Priority Queue. What exactly is different about their priority values?")
question("2. Which algorithm expanded fewer nodes in your mediumMaze experiment?")
question("3. Did the algorithm that expanded fewer nodes also return the lower-cost path?")
question("4. Why can UCS guarantee the least-cost solution with non-negative step costs while Greedy cannot?")
question("5. Which algorithm uses information about the goal location?")

# ---------------- Submission ----------------
heading("Submission and Reflection")
para("search.py and searchAgents.py exported from Colab; only those two files were edited.")
shot("h1_download")
question("What changed when search became informed? Which heuristic would you choose for this "
         "Pac-Man movement model, and why?")

os.makedirs(OUTDIR, exist_ok=True)
doc.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
